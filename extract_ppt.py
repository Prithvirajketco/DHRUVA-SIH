import pptx
import sys

def extract_text_from_shape(shape):
    text = []
    if shape.has_text_frame:
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                text.append(run.text)
    if shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                text.append(cell.text_frame.text)
    if shape.shape_type == 6: # Group shape
        for child in shape.shapes:
            text.extend(extract_text_from_shape(child))
    return text

def extract_text(filename):
    prs = pptx.Presentation(filename)
    text_runs = []
    for i, slide in enumerate(prs.slides):
        text_runs.append(f"\n--- [SLIDE {i+1}] ---")
        for shape in slide.shapes:
            text_runs.extend(extract_text_from_shape(shape))
        if slide.has_notes_slide:
            notes_frame = slide.notes_slide.notes_text_frame
            text_runs.append("\n[NOTES]:")
            text_runs.append(notes_frame.text)
    return "\n".join(text_runs)

if __name__ == "__main__":
    text = extract_text(sys.argv[1])
    with open("ppt_content_full.txt", "w", encoding="utf-8") as f:
        f.write(text)
    print("Extracted to ppt_content_full.txt")
